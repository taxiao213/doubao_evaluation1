#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""M4-01/02/03 华晟精密盈利预测模型（可复算）。
历史数据取自 PDF（单位万元，静态黑字）；预测驱动为蓝色输入假设；
预测与派生值全部为引用单元格的公式；跨 sheet 引用绿字。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/Users/secneo/Documents/workspace/ai_project/doubao_work2/测评执行/产物/M4_金融建模/华晟精密_盈利预测模型.xlsx"

FONT = "Calibri"
# 财务模型颜色编码
BLUE = Font(name=FONT, size=10, color="0000FF")          # 输入假设
BLACK = Font(name=FONT, size=10, color="000000")         # 本表公式/历史
BLACKB = Font(name=FONT, size=10, bold=True, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")         # 跨表引用
AUX = Font(name=FONT, size=10, italic=True, color="808080")
WHITE_B = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE_F = Font(name=FONT, size=12, bold=True, color="FFFFFF")
LBL = Font(name=FONT, size=10, color="000000")
LBL_B = Font(name=FONT, size=10, bold=True, color="000000")

HDR_BG = "0F3B5D"; SEC_BG = "3288B9"; LIGHT_BG = "CFE3F1"; BASE_BG = "FFF2CC"
fill_hdr = PatternFill("solid", fgColor=HDR_BG)
fill_sec = PatternFill("solid", fgColor=SEC_BG)
fill_light = PatternFill("solid", fgColor=LIGHT_BG)
fill_base = PatternFill("solid", fgColor=BASE_BG)

thin = Side(style="thin", color="808080")
bd_top = Border(top=thin)
bd_final = Border(top=thin, bottom=Side(style="double"))

NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%'
EPSF = '#,##0.00;(#,##0.00);"-"'
RGT = Alignment(horizontal="right", vertical="center")
LFT = Alignment(horizontal="left", vertical="center")
CTR = Alignment(horizontal="center", vertical="center")

wb = Workbook()

def title_bar(ws, text, c1, c2, row=1):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c); cell.fill = fill_hdr
    t = ws.cell(row=row, column=c1, value=text); t.font = TITLE_F; t.alignment = LFT
    ws.row_dimensions[row].height = 24

def section(ws, row, c1, c2, text):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).fill = fill_sec
    t = ws.cell(row=row, column=c1, value=text); t.font = WHITE_B; t.alignment = LFT

# ============ Sheet 1: Assumptions ============
a = wb.active; a.title = "Assumptions"
a.sheet_view.showGridLines = False
title_bar(a, "华晟精密 盈利预测模型 — 假设与驱动（Assumptions）", 2, 5)
a["B2"] = "单位：万元（除增速、比率与每股收益外）；蓝色字体为可修改的输入假设"
a["B2"].font = AUX
section(a, 4, 2, 5, "一、预测驱动假设（蓝色输入，可修改）")
hdr = ["驱动假设", "2025E", "2026E", "2027E"]
for j, h in enumerate(hdr, 2):
    c = a.cell(row=5, column=j, value=h); c.font = WHITE_B; c.fill = fill_light; c.font = BLACKB
    c.alignment = CTR if j > 2 else LFT
drivers = [
    ("营业收入增速", [0.15, 0.12, 0.10], PCT),
    ("归母净利率", [0.115, 0.118, 0.120], PCT),
    ("研发费用率（占营收）", [0.168, 0.165, 0.160], PCT),
    ("经营现金流净额/营收", [0.128, 0.130, 0.132], PCT),
]
r = 6
for name, vals, fmt in drivers:
    a.cell(row=r, column=2, value=name).font = LBL
    for j, v in enumerate(vals, 3):
        c = a.cell(row=r, column=j, value=v); c.font = BLUE; c.number_format = fmt; c.alignment = RGT
    r += 1
# 总股本：单一假设
a.cell(row=10, column=2, value="总股本（万股，按 2024 年 EPS 反推）").font = LBL
a.merge_cells("C10:E10")
c = a.cell(row=10, column=3, value=6309.18); c.font = BLUE; c.number_format = NUM; c.alignment = RGT
section(a, 12, 2, 5, "二、历史基准（2024A，取自年报 PDF，静态实际值）")
hist = [
    ("2024 年营业收入（万元，脚注①剔除一次性处置收益）", 53942),
    ("2024 年归母净利润（万元）", 6183),
    ("2024 年研发投入（万元，脚注③全部费用化）", 9051),
    ("2024 年经营活动现金流净额（万元）", 6892),
    ("2024 年基本每股收益（元）", 0.98),
]
r = 13
for name, v in hist:
    a.cell(row=r, column=2, value=name).font = LBL
    a.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    c = a.cell(row=r, column=3, value=v); c.font = BLACK; c.alignment = RGT
    c.number_format = EPSF if r == 17 else NUM
    r += 1
a["B19"] = "Methodology: 预测期营收=上年营收×(1+营收增速)；归母净利=营收×归母净利率；研发投入=营收×研发费用率；经营现金流=营收×CFO率；EPS=归母净利/总股本。"
a["B19"].font = AUX; a.merge_cells("B19:E20"); a["B19"].alignment = Alignment(wrap_text=True, vertical="top")
a.column_dimensions["A"].width = 2
a.column_dimensions["B"].width = 38
for col in "CDE": a.column_dimensions[col].width = 12

# ============ Sheet 2: 盈利预测 ============
m = wb.create_sheet("盈利预测")
m.sheet_view.showGridLines = False
title_bar(m, "华晟精密 盈利预测（2022A–2027E）", 2, 8)
m["B2"] = "单位：万元（除比率与每股收益）；历史 A 为实际值（黑字），预测 E 为公式；跨表引用绿字"
m["B2"].font = AUX
years = ["指标", "2022A", "2023A", "2024A", "2025E", "2026E", "2027E"]
for j, h in enumerate(years, 2):
    c = m.cell(row=4, column=j, value=h); c.font = WHITE_B; c.fill = fill_sec; c.alignment = CTR
m.cell(row=4, column=2).alignment = LFT
# 列字母 C..H = 2022..2027
cols = ["C", "D", "E", "F", "G", "H"]
# 行定义
# 5 营收  6 营收增速  7 归母净利  8 归母净利率  9 研发投入  10 研发费用率
# 11 经营现金流 12 CFO率 13 总股本 14 EPS
def put_label(r, text, bold=False):
    c = m.cell(row=r, column=2, value=text); c.font = LBL_B if bold else LBL; c.alignment = LFT
put_label(5, "营业收入（万元）")
put_label(6, "  营收增速")
put_label(7, "归母净利润（万元）", bold=True)
put_label(8, "  归母净利率")
put_label(9, "研发投入（万元）")
put_label(10, "  研发费用率")
put_label(11, "经营活动现金流净额（万元）")
put_label(12, "  CFO/营收")
put_label(13, "总股本（万股）")
put_label(14, "基本每股收益（元）", bold=True)

# 历史值（静态，黑字）
hist_rev = {"C": 38417, "D": 45166, "E": 53942}
hist_ni = {"C": 3860, "D": 4518, "E": 6183}
hist_rd = {"C": 6120, "D": 7483, "E": 9051}
hist_cfo = {"C": 4204, "D": 5377, "E": 6892}
hist_eps = {"C": 0.64, "D": 0.73, "E": 0.98}
for col, v in hist_rev.items():
    m[f"{col}5"] = v; m[f"{col}5"].number_format = NUM; m[f"{col}5"].font = BLACK; m[f"{col}5"].alignment = RGT
for col, v in hist_ni.items():
    m[f"{col}7"] = v; m[f"{col}7"].number_format = NUM; m[f"{col}7"].font = BLACKB; m[f"{col}7"].alignment = RGT
for col, v in hist_rd.items():
    m[f"{col}9"] = v; m[f"{col}9"].number_format = NUM; m[f"{col}9"].font = BLACK; m[f"{col}9"].alignment = RGT
for col, v in hist_cfo.items():
    m[f"{col}11"] = v; m[f"{col}11"].number_format = NUM; m[f"{col}11"].font = BLACK; m[f"{col}11"].alignment = RGT
for col, v in hist_eps.items():
    m[f"{col}14"] = v; m[f"{col}14"].number_format = EPSF; m[f"{col}14"].font = BLACK; m[f"{col}14"].alignment = RGT

# 历史比率/增速（本表公式，黑字）
for col, prev in (("D", "C"), ("E", "D")):
    m[f"{col}6"] = f"={col}5/{prev}5-1"
    m[f"{col}8"] = f"={col}7/{col}5"
    m[f"{col}10"] = f"={col}9/{col}5"
    m[f"{col}12"] = f"={col}11/{col}5"
    m[f"{col}13"] = f"={col}7/{col}14"
for addr in ["D6","E6","D8","E8","D10","E10","D12","E12"]:
    m[addr].number_format = PCT; m[addr].font = BLACK; m[addr].alignment = RGT
for addr in ["D13","E13"]:
    m[addr].number_format = NUM; m[addr].font = BLACK; m[addr].alignment = RGT
m["C13"] = "=C7/C14"; m["C13"].number_format = NUM; m["C13"].font = BLACK; m["C13"].alignment = RGT

# 预测列（跨表引用 Assumptions，绿字）
asm = {"F": "C", "G": "D", "H": "E"}  # 预测列 -> Assumptions 年度列
prev_col = {"F": "E", "G": "F", "H": "G"}
for fc, ac in asm.items():
    pc = prev_col[fc]
    # 营收
    m[f"{fc}5"] = f"={pc}5*(1+Assumptions!${ac}$6)"; m[f"{fc}5"].number_format = NUM
    # 营收增速 = 引用假设
    m[f"{fc}6"] = f"=Assumptions!${ac}$6"; m[f"{fc}6"].number_format = PCT
    # 归母净利
    m[f"{fc}7"] = f"={fc}5*Assumptions!${ac}$7"; m[f"{fc}7"].number_format = NUM
    # 净利率
    m[f"{fc}8"] = f"=Assumptions!${ac}$7"; m[f"{fc}8"].number_format = PCT
    # 研发
    m[f"{fc}9"] = f"={fc}5*Assumptions!${ac}$8"; m[f"{fc}9"].number_format = NUM
    m[f"{fc}10"] = f"=Assumptions!${ac}$8"; m[f"{fc}10"].number_format = PCT
    # CFO
    m[f"{fc}11"] = f"={fc}5*Assumptions!${ac}$9"; m[f"{fc}11"].number_format = NUM
    m[f"{fc}12"] = f"=Assumptions!${ac}$9"; m[f"{fc}12"].number_format = PCT
    # 股本
    m[f"{fc}13"] = "=Assumptions!$C$10"; m[f"{fc}13"].number_format = NUM
    # EPS
    m[f"{fc}14"] = f"={fc}7/Assumptions!$C$10"; m[f"{fc}14"].number_format = EPSF
    for rr in (5,6,7,8,9,10,11,12,13,14):
        cell = m[f"{fc}{rr}"]; cell.font = GREEN; cell.alignment = RGT
# 关键小计/合计边框
for rr in (7, 14):
    for cc in range(2, 9):
        m.cell(row=rr, column=cc).border = bd_top
m["B16"] = "Sources: 历史数据取自《华晟精密2024年报节选.pdf》第2页主要财务数据（虚构公司，供测评）；预测假设见 Assumptions 表。"
m["B16"].font = AUX; m.merge_cells("B16:H16")
m.column_dimensions["A"].width = 2
m.column_dimensions["B"].width = 30
for col in cols: m.column_dimensions[col].width = 11

# ============ Sheet 3: 敏感性分析（M4-03） ============
s = wb.create_sheet("敏感性分析")
s.sheet_view.showGridLines = False
title_bar(s, "敏感性分析：2025E 营收增速假设 ±5pct 全链路联动", 2, 5)
s["B2"] = "仅对 2025E 营收增速做 -5pct / 基准 / +5pct 情景，2026–2027 增速与各率维持基准；结果为全链路重算公式"
s["B2"].font = AUX; s.merge_cells("B2:F2")
for j, h in enumerate(["情景", "悲观（-5pct）", "基准", "乐观（+5pct）"], 2):
    c = s.cell(row=4, column=j, value=h); c.font = WHITE_B; c.fill = fill_sec; c.alignment = CTR
# 行：5 2025增速 6 2025营收 7 2026营收 8 2027营收 9 2027归母净利 10 2027 EPS
labels = ["2025E 营收增速", "2025E 营业收入（万元）", "2026E 营业收入（万元）",
          "2027E 营业收入（万元）", "2027E 归母净利润（万元）", "2027E 基本 EPS（元）"]
for i, t in enumerate(labels, 5):
    s.cell(row=i, column=2, value=t).font = LBL
scn_cols = {"C": "-0.05", "D": "+0", "E": "+0.05"}
for sc, delta in scn_cols.items():
    s[f"{sc}5"] = f"=Assumptions!$C$6{delta}"
    s[f"{sc}6"] = f"=盈利预测!$E$5*(1+{sc}5)"
    s[f"{sc}7"] = f"={sc}6*(1+Assumptions!$D$6)"
    s[f"{sc}8"] = f"={sc}7*(1+Assumptions!$E$6)"
    s[f"{sc}9"] = f"={sc}8*Assumptions!$E$7"
    s[f"{sc}10"] = f"={sc}9/Assumptions!$C$10"
    for rr in range(5, 11):
        cell = s[f"{sc}{rr}"]; cell.font = GREEN; cell.alignment = RGT
        cell.number_format = PCT if rr == 5 else (EPSF if rr == 10 else NUM)
# 基准列浅黄高亮
for rr in range(4, 11):
    s.cell(row=rr, column=4).fill = fill_base
    if rr >= 5:
        s.cell(row=rr, column=4).font = BLACKB
s["B12"] = "说明：修改 Assumptions 表中的营收增速或净利率，本表三种情景全部联动重算；基准列（浅黄）与主模型结果一致。"
s["B12"].font = AUX; s.merge_cells("B12:F12")
s.column_dimensions["A"].width = 2; s.column_dimensions["B"].width = 28
for col in "CDE": s.column_dimensions[col].width = 15

# ============ Sheet 4: 勾稽校验 ============
k = wb.create_sheet("勾稽校验")
k.sheet_view.showGridLines = False
title_bar(k, "PDF 数据勾稽与跨页引用校验", 2, 5)
k["B2"] = "单位：万元；差异=重算值-披露值，应为 0"
k["B2"].font = AUX
section(k, 4, 2, 5, "一、分季度数据加总 vs 全年披露")
for j, h in enumerate(["校验项", "Q1", "Q2", "Q3", "Q4"], 2):
    c = k.cell(row=5, column=j, value=h); c.font = BLACKB; c.fill = fill_light; c.alignment = CTR
# 分季度
q_rev = [12104, 13566, 14220, 14052]
q_ni = [1302, 1494, 1683, 1704]
k.cell(row=6, column=2, value="分季度营收").font = LBL
k.cell(row=7, column=2, value="分季度归母净利").font = LBL
for j, v in enumerate(q_rev, 3):
    cc = k.cell(row=6, column=j, value=v); cc.number_format = NUM; cc.alignment = RGT
for j, v in enumerate(q_ni, 3):
    cc = k.cell(row=7, column=j, value=v); cc.number_format = NUM; cc.alignment = RGT
# 研发构成（先落明细，供勾稽引用）
section(k, 9, 2, 6, "二、研发投入构成（附注3，单位万元）")
k.cell(row=10, column=2, value="研发投入构成").font = LBL
for j, h in enumerate(["人员人工", "材料试制", "折旧摊销", "其他"], 3):
    c = k.cell(row=10, column=j, value=h); c.font = BLACKB; c.fill = fill_light; c.alignment = CTR
for j, v in enumerate([4866, 2170, 1205, 810], 3):
    cc = k.cell(row=11, column=j, value=v); cc.number_format = NUM; cc.alignment = RGT
# 勾稽差异小表
section(k, 13, 2, 5, "三、勾稽差异（应为 0）")
for j, h in enumerate(["校验项", "重算值", "披露值", "差异"], 2):
    c = k.cell(row=14, column=j, value=h); c.font = BLACKB; c.fill = fill_light; c.alignment = CTR
rows = [
    ("分季度营收合计 = 全年营收", "=SUM(C6:F6)", 53942),
    ("分季度净利合计 = 全年归母净利", "=SUM(C7:F7)", 6183),
    ("研发投入构成合计 = 披露研发投入", "=SUM(C11:F11)", 9051),
]
r = 15
for name, formula, disc in rows:
    k.cell(row=r, column=2, value=name).font = LBL
    k.cell(row=r, column=3, value=formula).number_format = NUM
    k.cell(row=r, column=4, value=disc).number_format = NUM
    k.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = NUM
    for cc in range(3, 6): k.cell(row=r, column=cc).alignment = RGT
    r += 1
# 跨页引用
section(k, 19, 2, 6, "四、跨页引用")
k.cell(row=20, column=2, value="第1页脚注③ → 指向第2页「附注3：研发投入口径」（全部费用化、资本化率为0）。").font = BLACK
k.merge_cells("B20:F20")
k.cell(row=21, column=2, value="脚注①：营业收入为剔除一次性处置收益后口径；脚注②：归母净利润为归属于上市公司股东的净利润。").font = BLACK
k.merge_cells("B21:F21")
k.column_dimensions["A"].width = 2; k.column_dimensions["B"].width = 34
for col in "CDEF": k.column_dimensions[col].width = 12

# ============ Sheet 5: 可比公司（M4-02） ============
cmp = wb.create_sheet("可比公司")
cmp.sheet_view.showGridLines = False
title_bar(cmp, "可比公司对比（2024A，精密制造）", 2, 5)
cmp["B2"] = "立讯精密/领益智造单位为亿元（真实公开数据）；华晟精密为虚构样本，由万元换算（÷10000）；派生指标为公式"
cmp["B2"].font = AUX; cmp.merge_cells("B2:F2")
for j, h in enumerate(["指标", "立讯精密(002475)", "领益智造(002600)", "华晟精密(样本)"], 2):
    c = cmp.cell(row=4, column=j, value=h); c.font = WHITE_B; c.fill = fill_sec; c.alignment = CTR
# 行
# 5 2023营收 6 2024营收 7 2023归母 8 2024归母 9 营收增速 10 归母增速
# 11 2024毛利率 12 2024净利率 13 2024研发费用 14 研发费用率 15 2024 ROE
data = {
    5: ("2023 营业收入（亿元）", [2319.05, 341.24, 4.5166], NUM),
    6: ("2024 营业收入（亿元）", [2687.95, 442.60, 5.3942], NUM),
    7: ("2023 归母净利润（亿元）", [109.53, 20.51, 0.4518], NUM),
    8: ("2024 归母净利润（亿元）", [133.66, 17.55, 0.6183], NUM),
    11: ("2024 销售毛利率", [0.1041, 0.1577, None], PCT),
    13: ("2024 研发费用（亿元）", [85.56, 19.90, 0.9051], NUM),
    15: ("2024 净资产收益率 ROE", [0.1928, 0.0885, None], PCT),
}
for rr, (name, vals, fmt) in data.items():
    cmp.cell(row=rr, column=2, value=name).font = LBL
    for j, v in enumerate(vals, 3):
        c = cmp.cell(row=rr, column=j)
        if v is not None:
            c.value = v; c.number_format = fmt; c.font = BLACK
        else:
            c.value = "未披露"; c.font = AUX
        c.alignment = RGT
# 派生公式行
cmp.cell(row=9, column=2, value="  营业收入增速").font = LBL
cmp.cell(row=10, column=2, value="  归母净利润增速").font = LBL
cmp.cell(row=12, column=2, value="2024 归母净利率").font = LBL
cmp.cell(row=14, column=2, value="  研发费用率").font = LBL
for col in ("C", "D", "E"):
    cmp[f"{col}9"] = f"={col}6/{col}5-1"; cmp[f"{col}9"].number_format = PCT
    cmp[f"{col}10"] = f"={col}8/{col}7-1"; cmp[f"{col}10"].number_format = PCT
    cmp[f"{col}12"] = f"={col}8/{col}6"; cmp[f"{col}12"].number_format = PCT
    cmp[f"{col}14"] = f"={col}13/{col}6"; cmp[f"{col}14"].number_format = PCT
    for rr in (9, 10, 12, 14):
        cmp[f"{col}{rr}"].font = BLACK; cmp[f"{col}{rr}"].alignment = RGT
cmp["B17"] = "Sources: 立讯精密(002475)、领益智造(002600) 2023/2024 年报数据来自豆包金融数据（同花顺口径）；华晟精密为 PDF 虚构样本，毛利率、ROE 未在节选年报中披露。"
cmp["B17"].font = AUX; cmp.merge_cells("B17:F18"); cmp["B17"].alignment = Alignment(wrap_text=True, vertical="top")
cmp.column_dimensions["A"].width = 2; cmp.column_dimensions["B"].width = 28
for col in "CDE": cmp.column_dimensions[col].width = 18

# ============ Sheet 6: 来源与口径 ============
src = wb.create_sheet("来源与口径")
src.sheet_view.showGridLines = False
title_bar(src, "数据来源、口径与颜色约定", 2, 4)
notes = [
    ("历史数据", "华晟精密2024年报节选.pdf（虚构公司，证券代码839201，披露日期2025-03-28），第2页主要财务数据，单位万元。"),
    ("脚注①", "营业收入为剔除一次性处置收益后的口径。"),
    ("脚注②", "归母净利润指归属于上市公司股东的净利润。"),
    ("脚注③ / 跨页引用", "研发投入全部费用化、资本化率为0；第1页脚注③指向第2页「附注3：研发投入口径」。"),
    ("分季度勾稽", "分季度营收合计 53,942 = 全年营收；分季度净利合计 6,183 = 全年归母净利；研发构成合计 9,051 = 研发投入。"),
    ("可比公司", "立讯精密(002475)、领益智造(002600) 2023/2024 年报数据，来源豆包金融数据（同花顺利润业绩表口径），单位亿元。"),
    ("预测假设", "2025–2027 营收增速、净利率、研发费用率、CFO率为测评演示假设（蓝色），非券商预测，可在 Assumptions 表修改。"),
    ("颜色约定", "蓝色字体=可修改输入假设；黑色=历史实际值/本表公式；绿色=跨工作表引用；灰色斜体=说明注释。"),
]
r = 4
for k0, v in notes:
    src.cell(row=r, column=2, value=k0).font = LBL_B
    c = src.cell(row=r, column=3, value=v); c.font = BLACK; c.alignment = Alignment(wrap_text=True, vertical="top")
    src.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    src.row_dimensions[r].height = 30
    r += 1
src.column_dimensions["A"].width = 2; src.column_dimensions["B"].width = 18
for col in "CDEF": src.column_dimensions[col].width = 22

wb.save(OUT)
print("saved:", OUT)
