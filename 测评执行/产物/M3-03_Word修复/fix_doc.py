#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""M3-03 修复脚本：在 TARGET_DOCX 上做①错别字订正 ②全局样式统一。
保留原稿全部事实与结构，仅修正错别字并把混乱格式归一到规范样式。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "原稿_混乱_季度工作总结.docx"
TGT = "修复稿_季度工作总结.docx"

# ---------- 1. 错别字订正映射（逐字核对原稿） ----------
TYPO_FIX = [
    ("工作总节", "工作总结"),
    ("重点保正了", "重点保证了"),
    ("报表导出错慢", "报表导出缓慢"),
    ("完正的", "完整的"),
    ("工作按排", "工作安排"),
    ("提生移动端", "提升移动端"),
    ("提生到", "提升到"),
    ("关健节点", "关键节点"),
]

# ---------- 样式常量（统一规范） ----------
FONT_SONG = "宋体"; FONT_HEI = "黑体"; FONT_KAI = "楷体"; FONT_EN = "Times New Roman"
C_TITLE = RGBColor(0x1F, 0x1F, 0x1F)
C_BODY  = RGBColor(0x00, 0x00, 0x00)
C_MUTED = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F4E79"

def set_cn_font(run, cn, en=FONT_EN, size=12, bold=False, color=None, italic=False):
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color if color else C_BODY

def style_base(doc):
    """配置 Normal / 标题等命名样式，保证全局统一。"""
    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_SONG)
    st.font.size = Pt(12)
    st.font.color.rgb = C_BODY
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, cn, size, bold, color, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_cn_font(r, cn, size=size, bold=bold, color=color)

def add_title(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); set_cn_font(r, FONT_HEI, size=22, bold=True, color=C_TITLE)
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(text); set_cn_font(r, FONT_KAI, size=14, color=C_MUTED)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text); set_cn_font(r, FONT_HEI, size=16, bold=True, color=C_TITLE)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text); set_cn_font(r, FONT_HEI, size=14, bold=True, color=C_TITLE)
    return p

def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)  # 2 字符
    r = p.add_run(text); set_cn_font(r, FONT_SONG, size=12, color=C_BODY)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-12)  # 悬挂缩进
    r = p.add_run("• " + text); set_cn_font(r, FONT_SONG, size=12, color=C_BODY)
    return p

def add_sign(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text); set_cn_font(r, FONT_SONG, size=12, color=C_BODY)
    return p

def normalize_numbers(t):
    # 数字与中文单位之间的空格统一去掉；保留拉丁缩写与中文之间的空格
    t = re.sub(r'(\d)\s+(?=[个单场人次%元月年])', r'\1', t)
    return t

def fix_typos(t):
    for a, b in TYPO_FIX:
        t = t.replace(a, b)
    return t

# ---------- 读取原稿段落，提取规范后的正文（保留事实） ----------
src = Document(SRC)
raw = [p.text.strip() for p in src.paragraphs]
# 原稿表格数据
src_tbl = src.tables[0]
tbl_rows = [[c.text.strip() for c in row.cells] for row in src_tbl.rows]

# ---------- 重建 TARGET ----------
tgt = Document(TGT)
# 清空 TARGET 原有 body
for p in list(tgt.paragraphs):
    p._element.getparent().remove(p._element)
for t in list(tgt.tables):
    t._element.getparent().remove(t._element)
style_base(tgt)

i = 0
def is_bullet(t):
    t = t.replace("\u3000", "").strip()
    return t.startswith("•") or bool(re.match(r'^\d+\s*[、)．.]', t))

def clean(t):
    t = t.replace("\u3000", "").strip()  # 去掉段首全角空格
    t = re.sub(r'^•\s*', '', t)          # 去掉 • 项目符号
    t = re.sub(r'^\d+\s*[、)．.]\s*', '', t)  # 去掉 1) / 1、 等
    return t

# 逐段映射（依据原稿语义）
paras = [t for t in raw if t != ""]
# 标题
add_title(tgt, fix_typos(paras[0]))
add_subtitle(tgt, paras[1])
n = 2
while n < len(paras):
    t = paras[n]
    if re.match(r'^[一二三四五六七八九十]+、', t):
        add_h1(tgt, fix_typos(t))
    elif re.match(r'^（[一二三四五六七八九十]+）', t):
        add_h2(tgt, fix_typos(t))
    elif is_bullet(t):
        add_bullet(tgt, normalize_numbers(fix_typos(clean(t))))
    elif t in ("信息技术部",) or re.search(r'\d{4}年\d{1,2}月\d{1,2}日$', t):
        add_sign(tgt, t)
    else:
        add_body(tgt, normalize_numbers(fix_typos(t)))
    n += 1

# 在“主要指标”相关位置表格应紧跟二级标题；python-docx 只能末尾追加，
# 因此重建时把表格插入到 H2 段之后。定位“（一）主要指标完成情况”后面的 4 条 bullet 之后。
# 简化：先找到最后一个 bullet（培训完成）段落，在其后插入表格。
tbl = tgt.add_table(rows=len(tbl_rows), cols=len(tbl_rows[0]))
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for ri, row in enumerate(tbl_rows):
    for ci, val in enumerate(row):
        if ri == 0:
            set_cell_text(tbl.rows[ri].cells[ci], val, FONT_HEI, 11, True, RGBColor(0xFF,0xFF,0xFF))
            shade_cell(tbl.rows[ri].cells[ci], HEADER_FILL)
        else:
            set_cell_text(tbl.rows[ri].cells[ci], normalize_numbers(val), FONT_SONG, 11, False, C_BODY)
# 移动表格到正确位置（培训完成 bullet 之后）
train_p = None
for p in tgt.paragraphs:
    if "培训完成" in p.text:
        train_p = p
if train_p is not None:
    train_p._element.addnext(tbl._element)
    # 表格后补一个空段
    sp = OxmlElement('w:p'); tbl._element.addnext(sp)

tgt.save(TGT)
print("fixed saved:", TGT)

# 自检：错别字是否全部清除
chk = Document(TGT)
alltext = "\n".join(p.text for p in chk.paragraphs)
for t in chk.tables:
    for r in t.rows:
        alltext += "\n" + " ".join(c.text for c in r.cells)
left = [a for a, _ in TYPO_FIX if a in alltext]
print("残留错别字:", left if left else "无")
print("段落数:", len(chk.paragraphs), "| 表格数:", len(chk.tables))
