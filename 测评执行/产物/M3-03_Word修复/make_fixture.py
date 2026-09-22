#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""M3-03 测试夹具：构造一份含错别字、格式混乱的 docx（模拟用户提供的原稿）。
注意：本文件仅用于测评输入，不是最终产物。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

doc = Document()

def set_run(run, font_cn="宋体", size=12, bold=False, color=None, font_en="Times New Roman", italic=False):
    run.font.name = font_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def para(text, font_cn="宋体", size=12, bold=False, align=None, color=None,
         space_before=0, space_after=0, line=None, font_en="Times New Roman", indent_first=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line:
        pf.line_spacing = line
    if indent_first:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    set_run(r, font_cn=font_cn, size=size, bold=bold, color=color, font_en=font_en)
    return p

# —— 标题：未用标题样式，手动加粗、黑体、二号、居中 ——（标题字号还不统一）
para("信息技术部 2026年第二季度工作总节", font_cn="黑体", size=22, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
# 副标题：楷体、小三、居中、蓝色
para("（汇报人：王磊　二〇二六年七月）", font_cn="楷体", size=15,
     align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x1F, 0x49, 0x7D), space_after=18)

# —— 一级标题：正文样式手动加粗，字号忽大忽小 ——
para("一、本季度工作完成情况", font_cn="黑体", size=16, bold=True, space_before=10, space_after=6)
# 正文：宋体/仿宋混用，字号 12/10.5 混用，行距不一，段首缩进有的有有的没有
para("　　二季度，部门按年度计划推进各项任务，整体进展顺利。在系统运维方面，我们重点保正了核心业务系统的稳定运行，"
     "可用率达到99.95%，未发生一起一级故障。", font_cn="宋体", size=12, line=1.5, space_after=4)
para("在项目建设方面，新上线的客户管理系统（CRM）按计划完成了一期交付，覆盖销售、客服两个部门，"
     "用户反馈总体良好，但也存在报表导出错慢、部分页面加载时间偏长等问题，需要在下个阶段重点优化。",
     font_cn="仿宋", size=10.5, line=1.15, space_after=8)

# 二级标题：用了正文样式、手动加粗、颜色还是红色
para("（一）主要指标完成情况", font_cn="宋体", size=13, bold=True, color=(0xC0, 0x00, 0x00), space_before=4, space_after=4)

# 手动项目符号，符号混用，缩进混乱
para("• 系统可用率：99.95%，高于季度目标99.9%；", font_cn="宋体", size=12, line=1.4, space_after=2)
para("1) 需求交付：累计完成需求 87个，其中紧急需求 15个，按期交付率 94%；", font_cn="仿宋", size=11, line=1.4, space_after=2)
para("1、  故障处理：共处理工单 236 单，平均响应时间 18 分钟，平均解决时间 3.2 小时；",
     font_cn="宋体", size=11, line=1.4, space_after=2)
para("•培训完成：组织内部培训 6 场，累计参训 214 人次。", font_cn="楷体", size=12, line=1.6, space_after=8)

# 空行
para("", size=12)

# 一级标题
para("二、存在的主要问题", font_cn="黑体", size=15, bold=True, space_before=10, space_after=6)  # 字号比上面的16小
para("　　对照年初目标，仍存在以下不足：一是跨部门协作的沟通成本偏高，部分需求在评审阶段反复变更，"
     "影响了开发节奏；二是自动化测试覆盖率不足，目前仅为 62%，回归测试仍较多依赖人工；"
     "三是文档管理不够规范，部分系统缺少完正的操作手册和运维手册。",
     font_cn="宋体", size=12, line=1.5, space_after=8)

# 表格：手动建，表头/正文字体字号不统一
table = doc.add_table(rows=4, cols=3)
table.style = "Table Grid"
hdr = ["指标", "目标值", "实际值"]
for j, t in enumerate(hdr):
    cell = table.rows[0].cells[j]
    cell.text = ""
    r = cell.paragraphs[0].add_run(t)
    set_run(r, font_cn="黑体", size=10, bold=True)  # 表头字号偏小
data = [["系统可用率", "99.9%", "99.95%"], ["需求按期交付率", "95%", "94%"], ["自动化测试覆盖率", "70%", "62%"]]
for i, row in enumerate(data):
    for j, t in enumerate(row):
        cell = table.rows[i + 1].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(t)
        # 表格正文字体/字号混乱
        set_run(r, font_cn=("仿宋" if i % 2 else "宋体"), size=(11 if j == 0 else 10),
                bold=(j == 2 and i == 1), color=((0xC0,0,0) if j == 2 and i == 1 else None))

para("", size=10)
# 一级标题
para("三、下季度工作按排", font_cn="黑体", size=16, bold=True, space_before=10, space_after=6)
para("　　三季度将重点做好以下工作：第一，持续提升系统稳定性，完善监控告警体系，力争把平均故障恢复时间压缩到 2 小时以内；"
     "第二，推进 CRM 二期建设，补齐报表性能短板，提生移动端使用体验；",
     font_cn="宋体", size=12, line=1.5, space_after=4)
para("第三，加强自动化测试能力建设，目标把覆盖率从62%提生到75%；第四，规范文档管理，"
     "在三季度末完成核心系统操作手册和运维手册的补全工作。各小组要明确责任人，关健节点提前向部门汇报。",
     font_cn="仿宋", size=11, line=1.3, space_after=10)

# 结尾：右对齐、楷体、字号偏大
para("信息技术部", font_cn="楷体", size=14, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
para("2026年7月10日", font_cn="楷体", size=14, align=WD_ALIGN_PARAGRAPH.RIGHT)

out = "/Users/secneo/Documents/workspace/ai_project/doubao_work2/测评执行/产物/M3-03_Word修复/原稿_混乱_季度工作总结.docx"
doc.save(out)
print("saved:", out)
