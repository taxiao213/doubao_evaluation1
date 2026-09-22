#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把财务摘要分析报表按真实公式结果与条件格式渲染为 HTML（无 LibreOffice 时的视觉自检）。
数值为公式重算结果（已由飞书服务端重算确认）；条件格式按 xlsx 中写入的规则模拟。"""
from openpyxl import load_workbook

wb = load_workbook("财务摘要分析报表.xlsx")
ws = wb["财务分析报表"]

# 真实重算结果（飞书回读确认）
calc = {
    "E4":0.194,"E5":0.173,"E6":1.0,"E7":0.210,"E8":0.367,"E9":-1.4,
    "B10":1627,"C10":1873,"D10":2293,"E10":0.224,
    "B11":0.1593,"C11":0.1656,"D11":0.1678,"E11":0.2,
    "B12":0.10047,"C12":0.10007,"D12":0.11457,"E12":1.5,
    "B14":1456,"B15":13753,
}
def disp(addr):
    c = ws[addr]; v = calc.get(addr, c.value)
    if isinstance(v, str) and v.startswith("="): return ""
    if v is None: return ""
    fmt = c.number_format
    if "pct" in fmt:
        return f"{v:+.1f}pct"
    if "%" in fmt:
        return f"{v*100:.1f}%" if addr not in ("E4","E5","E7","E8","E10") else f"{v*100:+.1f}%"
    if "#,##0" in fmt:
        return f"{v:,.0f}"
    return str(v)

def cell_color(addr, v):
    # 条件格式
    if addr in [f"E{r}" for r in range(4,13)]:
        if isinstance(v,(int,float)):
            if v > 0: return "background:#C6EFCE;color:#006100;"
            if v < 0: return "background:#FFC7CE;color:#9C0006;"
    if addr in ("B9","C9","D9") and isinstance(v,(int,float)) and v > 0.4:
        return "background:#FFC7CE;color:#9C0006;"
    if addr in ("B6","C6","D6") and isinstance(v,(int,float)) and v < 0.4:
        return "background:#FFF2CC;"
    return ""

# 数据条（金额行 2022-2024）
databar_rows = {4,5,7,8,10}
def bar_style(addr, raw):
    import re
    m = re.match(r"([A-Z])(\d+)", addr)
    col, row = m.group(1), int(m.group(2))
    if row in databar_rows and col in ("B","C","D") and isinstance(raw,(int,float)):
        # 同口径按全表金额最大值归一
        mx = 5394.0
        w = max(2, raw/mx*100)
        return f"background:linear-gradient(to right,#5B9BD5 {w:.0f}%,#fff {w:.0f}%);"
    return ""

rows_html = []
for r in range(3, 16):
    cells = []
    for ci, col in enumerate("ABCDE"):
        addr = f"{col}{r}"
        c = ws[addr]
        raw = calc.get(addr, c.value)
        if r == 3:
            cells.append(f'<th style="background:#5B9BD5;color:#fff;border-bottom:2px solid #17365D;">{c.value or ""}</th>')
            continue
        label = c.value
        if ci == 0:
            txt = label if isinstance(label,str) else ""
            italic = "font-style:italic;color:#1F4E79;" if r in (10,11,12) else "font-weight:bold;"
            top = "border-top:2px solid #17365D;" if r in (14,15) else ""
            cells.append(f'<td style="text-align:left;{italic}{top}">{txt}</td>')
        else:
            if r == 13:
                cells.append("<td></td>"); continue
            shown = disp(addr)
            style = "text-align:right;padding-right:8px;"
            style += cell_color(addr, raw)
            style += bar_style(addr, raw if isinstance(raw,(int,float)) else None)
            if r in (14,15):
                style += "border-top:2px solid #17365D;font-weight:bold;"
            if r in (10,11,12):
                style += "font-style:italic;color:#1F4E79;"
            cells.append(f'<td style="{style}">{shown}</td>')
    rows_html.append("<tr>" + "".join(cells) + "</tr>")

html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
body{{font-family:'微软雅黑',sans-serif;background:#f0f0f0;margin:0;padding:30px;}}
.card{{background:#fff;width:760px;margin:0 auto;padding:30px 36px;box-shadow:0 2px 12px rgba(0,0,0,.15);}}
.title{{background:#17365D;color:#fff;font-size:16px;font-weight:bold;text-align:center;padding:10px;margin:0;}}
.sub{{font-size:10px;color:#595959;padding:6px 2px 12px;}}
table{{border-collapse:collapse;width:100%;font-size:11px;}}
th,td{{border:1px solid #B8C7D9;padding:6px 6px;}}
th{{padding:8px 6px;}}
td{{font-family:Arial;}}
</style></head><body><div class="card">
<div class="title">华晟精密财务摘要分析报表（2022–2024 年）</div>
<div class="sub">金额单位：万元；比率为百分比；数据来源：table1_财务摘要（CSV 与表格截图一致）</div>
<table>{''.join(rows_html)}</table>
</div></body></html>"""
open("M3-04_报表预览.html","w",encoding="utf-8").write(html)
print("saved M3-04_报表预览.html")
